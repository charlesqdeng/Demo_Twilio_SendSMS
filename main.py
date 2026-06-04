"""
Twilio SMS Bulk Messaging Application

This application sends SMS messages to recipients listed in an Excel file
using Twilio's Programmable Messaging API.
"""

import argparse
import logging
import re
import sys

from twilio.http.http_client import TwilioHttpClient
from twilio.rest import Client
from twilio.base.exceptions import TwilioRestException
from openpyxl import load_workbook
from pathlib import Path
from config import get_config

# Configure logging
logging.basicConfig(
    level=logging.INFO,
    format='%(asctime)s - %(levelname)s - %(message)s'
)
logger = logging.getLogger(__name__)


class ErrorCode:
    """Error codes for the application."""
    SUCCESS = "SUCCESS"
    INVALID_PHONE = "INVALID_PHONE"
    LOOKUP_FAILED = "LOOKUP_FAILED"
    SEND_FAILED = "SEND_FAILED"
    MISSING_DATA = "MISSING_DATA"
    INVALID_TEMPLATE = "INVALID_TEMPLATE"


class PhoneValidator:
    """Validate phone numbers using Twilio Lookup API."""

    def __init__(self, twilio_client):
        """Initialize phone validator with Twilio client."""
        self.client = twilio_client

    def validate_and_normalize(self, phone_number):
        """
        Validate phone number format and normalize it.
        
        Args:
            phone_number (str): Phone number to validate
            
        Returns:
            tuple: (is_valid, normalized_number, error_code)
        """
        if not phone_number:
            return False, None, ErrorCode.MISSING_DATA

        # Clean the phone number (remove common formatting)
        cleaned = re.sub(r'[\s\-\(\)\+]', '', str(phone_number).strip())

        # Check if it's numeric
        if not cleaned.isdigit():
            logger.warning(f"Non-numeric phone number: {phone_number}")
            return False, None, ErrorCode.INVALID_PHONE

        # Normalize: add leading 1 if 10 digits
        if len(cleaned) == 10:
            normalized = f"+1{cleaned}"
        elif len(cleaned) == 11 and cleaned.startswith('1'):
            normalized = f"+{cleaned}"
        else:
            logger.warning(f"Invalid phone number length: {phone_number} (length: {len(cleaned)})")
            return False, None, ErrorCode.INVALID_PHONE

        # Validate using Twilio Lookup API
        try:
            phone_data = self.client.lookups.v1.phone_numbers(normalized).fetch()
            logger.info(f"Phone validated: {normalized}")
            return True, normalized, None
        except TwilioRestException as e:
            logger.warning(f"Lookup failed for {phone_number}: {e.msg}")
            return False, None, ErrorCode.LOOKUP_FAILED
        except Exception as e:
            logger.error(f"Unexpected error validating {phone_number}: {str(e)}")
            return False, None, ErrorCode.LOOKUP_FAILED

    def format_for_display(self, phone_number):
        """Format validated phone number for display."""
        if phone_number:
            # Convert +1XXXXXXXXXX to (XXX) XXX-XXXX format
            digits = phone_number.replace('+', '')
            if len(digits) == 11 and digits.startswith('1'):
                digits = digits[1:]  # Remove leading 1
            if len(digits) == 10:
                return f"({digits[:3]}) {digits[3:6]}-{digits[6:]}"
        return phone_number


class SMSMessenger:
    """Handle SMS sending via Twilio."""

    def __init__(self, twilio_client, from_number, dry_run=False):
        """Initialize SMS messenger."""
        self.client = twilio_client
        self.from_number = from_number
        self.dry_run = dry_run

    def send_sms(self, to_number, message_body):
        """
        Send SMS message.
        
        Args:
            to_number (str): Recipient phone number
            message_body (str): Message content
            
        Returns:
            tuple: (success, error_code, message_sid)
        """
        if self.dry_run:
            logger.info(f"(dry-run) Would send SMS to {to_number}")
            return True, ErrorCode.SUCCESS, None
        try:
            message = self.client.messages.create(
                body=message_body,
                from_=self.from_number,
                to=to_number
            )
            logger.info(f"SMS sent successfully to {to_number} (SID: {message.sid})")
            return True, ErrorCode.SUCCESS, message.sid
        except TwilioRestException as e:
            logger.error(f"Failed to send SMS to {to_number}: {e.msg}")
            return False, ErrorCode.SEND_FAILED, None
        except Exception as e:
            logger.error(f"Unexpected error sending SMS to {to_number}: {str(e)}")
            return False, ErrorCode.SEND_FAILED, None


class ExcelProcessor:
    """Handle reading and writing Excel files."""

    def __init__(self, file_path):
        """Initialize Excel processor."""
        self.file_path = Path(file_path)
        if not self.file_path.exists():
            raise FileNotFoundError(f"Recipients file not found: {self.file_path}")

    def read_recipients(self):
        """
        Read recipient data from Excel file.
        
        Returns:
            list: List of recipient dictionaries with keys: name, phone, product_info
        """
        workbook = load_workbook(self.file_path)
        worksheet = workbook.active

        recipients = []
        for row_idx, row in enumerate(worksheet.iter_rows(values_only=False), start=1):
            # Skip header row (assuming row 1 is header)
            if row_idx == 1:
                continue

            # Extract data from columns 1-3
            try:
                name = row[0].value if row[0] else None
                phone = row[1].value if row[1] else None
                product_info = row[2].value if row[2] else None

                # Skip empty rows
                if not name or phone is None:
                    continue

                # Normalize phone values read from Excel, especially numeric cells
                if isinstance(phone, float):
                    phone = str(int(phone))
                elif isinstance(phone, int):
                    phone = str(phone)
                else:
                    phone = str(phone).strip()

                recipients.append({
                    'row_idx': row_idx,
                    'name': str(name).strip(),
                    'phone': phone,
                    'product_info': str(product_info).strip() if product_info else 'unknown',
                    'cell_ref': row[3]  # Reference to column 4 for error code
                })
            except Exception as e:
                logger.error(f"Error reading row {row_idx}: {str(e)}")
                continue

        logger.info(f"Read {len(recipients)} recipients from Excel file")
        return recipients

    def write_error_code(self, error_code):
        """
        Write error code to Excel file (Column 4).
        
        Args:
            error_code (dict): Dictionary with 'row_idx' and 'code'
        """
        workbook = load_workbook(self.file_path)
        worksheet = workbook.active

        row_idx = error_code['row_idx']
        code = error_code['code']

        # Write to column 4 (index 3)
        worksheet.cell(row=row_idx, column=4, value=code)

        workbook.save(self.file_path)

    def write_all_error_codes(self, error_codes):
        """
        Write all error codes to Excel file efficiently.
        
        Args:
            error_codes (list): List of dicts with 'row_idx' and 'code'
        """
        workbook = load_workbook(self.file_path)
        worksheet = workbook.active

        for error_code in error_codes:
            row_idx = error_code['row_idx']
            code = error_code['code']
            worksheet.cell(row=row_idx, column=4, value=code)

        workbook.save(self.file_path)
        logger.info(f"Updated {len(error_codes)} rows in Excel file")


class TemplateProcessor:
    """Handle message template processing."""

    @staticmethod
    def render(template, context):
        """
        Render message template with context variables.
        
        Args:
            template (str): Template string with {variable} placeholders
            context (dict): Dictionary with variable values
            
        Returns:
            tuple: (success, message, error_code)
        """
        try:
            # Normalize common template variable forms and render
            template = template.replace('{product info}', '{product_info}')
            formatted_context = {
                'name': context.get('name', 'Valued Customer'),
                'product_info': context.get('product_info', 'product'),
            }

            message = template.format(**formatted_context)
            return True, message, None
        except KeyError as e:
            logger.error(f"Template variable not found: {str(e)}")
            return False, None, ErrorCode.INVALID_TEMPLATE
        except Exception as e:
            logger.error(f"Error rendering template: {str(e)}")
            return False, None, ErrorCode.INVALID_TEMPLATE


class TwilioSMSApp:
    """Main application for sending SMS via Twilio."""

    def __init__(self, dry_run=False, recipients_file=None, template_file=None, from_number=None):
        """Initialize the application."""
        try:
            self.dry_run = dry_run
            self.config = get_config()

            if recipients_file:
                self.config.recipients_file = recipients_file
            if template_file:
                self.config.template_file = template_file
            if from_number:
                self.config.from_number = from_number

            self.twilio_client = Client(
                self.config.twilio_account_sid,
                self.config.twilio_auth_token,
                http_client=TwilioHttpClient(timeout=10)
            )
            self.phone_validator = PhoneValidator(self.twilio_client)
            self.messenger = SMSMessenger(self.twilio_client, self.config.from_number, dry_run=dry_run)
            self.excel_processor = ExcelProcessor(self.config.recipients_file)
            self.template = self.config.get_template()
            
            self.stats = {
                'total': 0,
                'sent': 0,
                'failed': 0,
                'invalid_phone': 0,
                'lookup_failed': 0,
                'send_failed': 0
            }
        except Exception as e:
            logger.error(f"Failed to initialize application: {str(e)}")
            raise

    def run(self):
        """Run the main application."""
        logger.info("Starting Twilio SMS Bulk Messaging Application")
        logger.info(f"From Number: {self.config.from_number}")
        logger.info(f"Recipients File: {self.config.recipients_file}")

        try:
            # Read recipients
            recipients = self.excel_processor.read_recipients()
            self.stats['total'] = len(recipients)

            error_codes = []

            # Process each recipient
            for recipient in recipients:
                logger.info(f"\nProcessing: {recipient['name']} ({recipient['phone']})")

                # Validate phone number
                is_valid, normalized_phone, error_code = self.phone_validator.validate_and_normalize(
                    recipient['phone']
                )

                if not is_valid:
                    logger.warning(f"Phone validation failed: {error_code}")
                    error_codes.append({
                        'row_idx': recipient['row_idx'],
                        'code': error_code
                    })
                    self.stats['failed'] += 1
                    if error_code == ErrorCode.INVALID_PHONE:
                        self.stats['invalid_phone'] += 1
                    elif error_code == ErrorCode.LOOKUP_FAILED:
                        self.stats['lookup_failed'] += 1
                    continue

                # Render message template
                success, message_body, error_code = TemplateProcessor.render(
                    self.template,
                    {
                        'name': recipient['name'],
                        'product_info': recipient['product_info']
                    }
                )

                if not success:
                    logger.error(f"Template rendering failed: {error_code}")
                    error_codes.append({
                        'row_idx': recipient['row_idx'],
                        'code': error_code
                    })
                    self.stats['failed'] += 1
                    continue

                # Send SMS
                success, error_code, message_sid = self.messenger.send_sms(
                    normalized_phone,
                    message_body
                )

                if not success:
                    logger.error(f"SMS send failed: {error_code}")
                    error_codes.append({
                        'row_idx': recipient['row_idx'],
                        'code': error_code
                    })
                    self.stats['failed'] += 1
                    self.stats['send_failed'] += 1
                else:
                    logger.info(f"✓ Message sent to {recipient['name']}")
                    error_codes.append({
                        'row_idx': recipient['row_idx'],
                        'code': ErrorCode.SUCCESS
                    })
                    self.stats['sent'] += 1

            # Write results to Excel unless running in dry-run mode
            if self.dry_run:
                logger.info("Dry-run mode enabled; no Excel changes were saved.")
            else:
                self.excel_processor.write_all_error_codes(error_codes)

            # Print summary
            self._print_summary()

        except Exception as e:
            logger.error(f"Application error: {str(e)}", exc_info=True)
            raise

    def _print_summary(self):
        """Print execution summary."""
        logger.info("\n" + "="*50)
        logger.info("EXECUTION SUMMARY")
        logger.info("="*50)
        logger.info(f"Total Recipients:     {self.stats['total']}")
        logger.info(f"Successfully Sent:    {self.stats['sent']}")
        logger.info(f"Failed:               {self.stats['failed']}")
        logger.info(f"  - Invalid Phone:    {self.stats['invalid_phone']}")
        logger.info(f"  - Lookup Failed:    {self.stats['lookup_failed']}")
        logger.info(f"  - Send Failed:      {self.stats['send_failed']}")
        logger.info("="*50)


def main():
    """Entry point for the application."""
    parser = argparse.ArgumentParser(description="Twilio SMS Bulk Messaging Application")
    parser.add_argument(
        '--dry-run',
        action='store_true',
        help='Validate recipients and template without sending SMS.'
    )
    parser.add_argument(
        '--recipients-file',
        help='Path to the recipients Excel file. Overrides RECIPIENTS_FILE in .env.'
    )
    parser.add_argument(
        '--template-file',
        help='Path to the message template file. Overrides TEMPLATE_FILE in .env.'
    )
    parser.add_argument(
        '--from-number',
        help='Twilio FROM_NUMBER override. Overrides FROM_NUMBER in .env.'
    )
    args = parser.parse_args()

    try:
        app = TwilioSMSApp(
            dry_run=args.dry_run,
            recipients_file=args.recipients_file,
            template_file=args.template_file,
            from_number=args.from_number
        )
        app.run()
        logger.info("Application completed successfully")
    except KeyboardInterrupt:
        logger.info("Application interrupted by user")
        sys.exit(1)
    except Exception as e:
        logger.error(f"Fatal error: {str(e)}")
        sys.exit(1)


if __name__ == '__main__':
    main()
